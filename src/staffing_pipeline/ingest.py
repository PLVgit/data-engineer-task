from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Iterable, Sequence

import duckdb
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet


class IngestionError(ValueError):
    """Raised when an input export violates its structural contract."""


@dataclass(frozen=True)
class SourceSpec:
    filename: str
    sheet_name: str
    table_name: str
    snapshot_label: str
    snapshot_date_format: str
    columns: tuple[tuple[str, str], ...]

    @property
    def source_headers(self) -> tuple[str, ...]:
        return tuple(source for source, _ in self.columns)

    @property
    def canonical_headers(self) -> tuple[str, ...]:
        return tuple(canonical for _, canonical in self.columns)


EMPLOYEES = SourceSpec(
    filename="hr_employees_export.xlsx",
    sheet_name="Employee Master Data",
    table_name="hr_employees",
    snapshot_label="Generated:",
    snapshot_date_format="%Y-%m-%d",
    columns=(
        ("Employee ID", "employee_id"),
        ("First Name", "first_name"),
        ("Last Name", "last_name"),
        ("Email Address", "email_address"),
        ("Department", "department"),
        ("Job Title", "job_title"),
        ("Date of Hire", "date_of_hire"),
        ("Termination Date", "termination_date"),
        ("Status", "status"),
        ("Reports To", "reports_to_employee_id"),
    ),
)

ASSIGNMENTS = SourceSpec(
    filename="project_assignments_report.xlsx",
    sheet_name="Project Assignments",
    table_name="project_assignments",
    snapshot_label="Report Date:",
    snapshot_date_format="%d/%m/%Y",
    columns=(
        ("Assignment ID", "assignment_id"),
        ("Emp. ID", "employee_id"),
        ("Project Code", "project_code"),
        ("Project Name", "project_name"),
        ("Assignment Role", "assignment_role"),
        ("Start Date", "start_date"),
        ("Weekly Hours", "weekly_hours"),
        ("Billable?", "billable_raw"),
    ),
)

SOURCE_SPECS = (EMPLOYEES, ASSIGNMENTS)
HEADER_SEARCH_ROWS = 10


def _normalize_header(value: object) -> str | None:
    if value is None:
        return None
    normalized = str(value).strip()
    return normalized or None


def _find_header_row(
    worksheet: Worksheet,
    required_headers: Sequence[str],
    search_rows: int = HEADER_SEARCH_ROWS,
) -> tuple[int, list[str | None]]:
    required = set(required_headers)
    candidates: list[tuple[int, int, list[str | None]]] = []

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=search_rows, values_only=True),
        start=1,
    ):
        headers = [_normalize_header(value) for value in row]
        overlap = len(required.intersection(header for header in headers if header))
        candidates.append((overlap, row_number, headers))

    if not candidates or max(item[0] for item in candidates) == 0:
        raise IngestionError(
            f"Could not identify a header row in the first {search_rows} rows of "
            f"sheet '{worksheet.title}'."
        )

    _, row_number, headers = max(candidates, key=lambda item: (item[0], -item[1]))
    return row_number, headers


def _validate_headers(
    headers: Sequence[str | None],
    required_headers: Sequence[str],
    *,
    filename: str,
    row_number: int,
) -> None:
    populated_headers = [header for header in headers if header is not None]
    counts = Counter(populated_headers)
    duplicate_required = sorted(
        header for header in required_headers if counts[header] > 1
    )
    missing = sorted(set(required_headers) - set(populated_headers))
    unexpected = sorted(set(populated_headers) - set(required_headers))

    problems: list[str] = []
    if missing:
        problems.append(f"missing required columns: {', '.join(missing)}")
    if duplicate_required:
        problems.append(
            f"duplicate required columns: {', '.join(duplicate_required)}"
        )
    if unexpected:
        problems.append(f"unexpected columns: {', '.join(unexpected)}")

    if problems:
        raise IngestionError(
            f"Invalid header row {row_number} in '{filename}': " + "; ".join(problems)
        )


def _read_snapshot_date(
    worksheet: Worksheet, spec: SourceSpec, header_row: int
) -> date:
    for row_number, row in enumerate(
        worksheet.iter_rows(max_row=header_row, values_only=True), start=1
    ):
        if row_number == header_row:
            break
        for value in row:
            if not isinstance(value, str):
                continue
            metadata = value.strip()
            if not metadata.startswith(spec.snapshot_label):
                continue
            raw_date = metadata.removeprefix(spec.snapshot_label).strip()
            try:
                return datetime.strptime(raw_date, spec.snapshot_date_format).date()
            except ValueError as exc:
                raise IngestionError(
                    f"Invalid snapshot date for '{spec.snapshot_label}' in "
                    f"'{spec.filename}': {raw_date!r}; "
                    f"expected {spec.snapshot_date_format}."
                ) from exc

    raise IngestionError(
        f"Missing required snapshot metadata '{spec.snapshot_label}' "
        f"before the header in '{spec.filename}'."
    )


def _raw_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    return str(value)


def read_source(path: Path, spec: SourceSpec) -> list[dict[str, str | int | date | None]]:
    if not path.is_file():
        raise IngestionError(f"Expected source file does not exist: '{path}'.")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise IngestionError(f"Could not open Excel export '{path}': {exc}") from exc

    try:
        if spec.sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames) or "none"
            raise IngestionError(
                f"Expected sheet '{spec.sheet_name}' in '{path.name}'; "
                f"available sheets: {available}."
            )

        worksheet = workbook[spec.sheet_name]
        header_row, headers = _find_header_row(worksheet, spec.source_headers)
        _validate_headers(
            headers,
            spec.source_headers,
            filename=path.name,
            row_number=header_row,
        )
        snapshot_date = _read_snapshot_date(worksheet, spec, header_row)

        source_positions = {
            header: index for index, header in enumerate(headers) if header is not None
        }
        rows: list[dict[str, str | int | date | None]] = []

        for source_row_number, values in enumerate(
            worksheet.iter_rows(min_row=header_row + 1, values_only=True),
            start=header_row + 1,
        ):
            selected_values = [
                values[source_positions[source_header]]
                if source_positions[source_header] < len(values)
                else None
                for source_header in spec.source_headers
            ]
            if all(value is None for value in selected_values):
                continue

            row: dict[str, str | int | date | None] = {
                canonical_header: _raw_text(value)
                for canonical_header, value in zip(
                    spec.canonical_headers, selected_values, strict=True
                )
            }
            row["_source_file"] = path.name
            row["_source_row_number"] = source_row_number
            row["_source_snapshot_date"] = snapshot_date
            rows.append(row)

        return rows
    finally:
        workbook.close()


def _replace_raw_table(
    connection: duckdb.DuckDBPyConnection,
    spec: SourceSpec,
    rows: Iterable[dict[str, str | int | date | None]],
) -> int:
    materialized_rows = list(rows)
    business_columns = spec.canonical_headers
    column_definitions = ", ".join(
        [
            *(f'"{column}" VARCHAR' for column in business_columns),
            '"_source_file" VARCHAR',
            '"_source_row_number" BIGINT',
            '"_source_snapshot_date" DATE',
        ]
    )
    qualified_table = f'raw."{spec.table_name}"'

    connection.execute(
        f"CREATE OR REPLACE TABLE {qualified_table} ({column_definitions})"
    )

    if materialized_rows:
        insert_columns = (
            *business_columns,
            "_source_file",
            "_source_row_number",
            "_source_snapshot_date",
        )
        placeholders = ", ".join("?" for _ in insert_columns)
        connection.executemany(
            f"INSERT INTO {qualified_table} VALUES ({placeholders})",
            [tuple(row[column] for column in insert_columns) for row in materialized_rows],
        )

    return len(materialized_rows)


def ingest(data_dir: Path, database_path: Path) -> dict[str, int]:
    extracted = {
        spec.table_name: read_source(data_dir / spec.filename, spec)
        for spec in SOURCE_SPECS
    }

    database_path.parent.mkdir(parents=True, exist_ok=True)
    connection = duckdb.connect(str(database_path))
    try:
        connection.execute("BEGIN TRANSACTION")
        connection.execute("CREATE SCHEMA IF NOT EXISTS raw")
        counts = {
            spec.table_name: _replace_raw_table(
                connection, spec, extracted[spec.table_name]
            )
            for spec in SOURCE_SPECS
        }
        connection.execute("COMMIT")
        return counts
    except Exception:
        connection.execute("ROLLBACK")
        raise
    finally:
        connection.close()


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Validate Excel exports and load canonical raw tables into DuckDB."
    )
    parser.add_argument("--data-dir", type=Path, default=Path("data"))
    parser.add_argument(
        "--database", type=Path, default=Path("warehouse/staffing.duckdb")
    )
    return parser.parse_args()


def main() -> None:
    args = _parse_args()
    counts = ingest(args.data_dir, args.database)
    for table_name, row_count in counts.items():
        print(f"Loaded {row_count} rows into raw.{table_name}")


if __name__ == "__main__":
    main()
