from dataclasses import replace
from datetime import date
from pathlib import Path

import duckdb
import pytest
from openpyxl import Workbook, load_workbook

from staffing_pipeline.ingest import (
    ASSIGNMENTS,
    EMPLOYEES,
    IngestionError,
    SourceSpec,
    ingest,
    read_source,
)


REPOSITORY_ROOT = Path(__file__).resolve().parents[1]


def _write_employee_workbook(
    path: Path,
    headers: list[str],
    *,
    sheet_name: str = EMPLOYEES.sheet_name,
    header_row: int = 4,
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = sheet_name
    worksheet.cell(row=1, column=1, value="Employee export")
    worksheet.cell(row=2, column=1, value="Generated: 2026-02-02")
    for column, header in enumerate(headers, start=1):
        worksheet.cell(row=header_row, column=column, value=header)

    values = {
        "Employee ID": "EMP-1",
        "First Name": "Ada",
        "Last Name": "Lovelace",
        "Email Address": "ada@example.com",
        "Department": "Engineering",
        "Job Title": "Engineer",
        "Date of Hire": "2024-01-01",
        "Termination Date": None,
        "Status": "Active",
        "Reports To": None,
    }
    for column, header in enumerate(headers, start=1):
        worksheet.cell(row=header_row + 1, column=column, value=values.get(header))
    workbook.save(path)


def test_detects_header_and_maps_canonical_columns(tmp_path: Path) -> None:
    source_path = tmp_path / EMPLOYEES.filename
    _write_employee_workbook(
        source_path,
        list(EMPLOYEES.source_headers),
        header_row=7,
    )

    rows = read_source(source_path, EMPLOYEES)

    assert len(rows) == 1
    assert rows[0]["employee_id"] == "EMP-1"
    assert rows[0]["email_address"] == "ada@example.com"
    assert rows[0]["_source_file"] == EMPLOYEES.filename
    assert rows[0]["_source_row_number"] == 8
    assert rows[0]["_source_snapshot_date"] == date(2026, 2, 2)
    assert set(rows[0]) == {
        *EMPLOYEES.canonical_headers,
        "_source_file",
        "_source_row_number",
        "_source_snapshot_date",
    }


def _write_snapshot_workbook(
    path: Path, spec: SourceSpec, metadata: str | None
) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = spec.sheet_name
    worksheet["A1"] = "Unrelated presentation text"
    worksheet["B2"] = metadata
    worksheet.append(spec.source_headers)
    worksheet.append(["ROW-1"])
    workbook.save(path)
    workbook.close()


@pytest.mark.parametrize(
    "spec, metadata",
    [
        (EMPLOYEES, "Generated: 2026-03-04"),
        (ASSIGNMENTS, "Report Date: 04/03/2026"),
    ],
)
def test_discovers_snapshot_date(
    tmp_path: Path, spec: SourceSpec, metadata: str
) -> None:
    source_path = tmp_path / spec.filename
    _write_snapshot_workbook(source_path, spec, metadata)

    rows = read_source(source_path, spec)

    assert len(rows) == 1
    assert rows[0]["_source_snapshot_date"] == date(2026, 3, 4)


@pytest.mark.parametrize(
    "spec, metadata, message",
    [
        (EMPLOYEES, None, "Missing required snapshot metadata"),
        (ASSIGNMENTS, None, "Missing required snapshot metadata"),
        (EMPLOYEES, "Generated: 2026-02-30", "Invalid snapshot date"),
        (ASSIGNMENTS, "Report Date: 30/02/2026", "Invalid snapshot date"),
    ],
)
def test_rejects_missing_or_malformed_snapshot_date(
    tmp_path: Path, spec: SourceSpec, metadata: str | None, message: str
) -> None:
    source_path = tmp_path / spec.filename
    _write_snapshot_workbook(source_path, spec, metadata)

    with pytest.raises(IngestionError, match=message) as error:
        read_source(source_path, spec)

    assert spec.filename in str(error.value)
    assert spec.snapshot_label in str(error.value)


def test_rejects_missing_required_column(tmp_path: Path) -> None:
    source_path = tmp_path / EMPLOYEES.filename
    headers = [
        header for header in EMPLOYEES.source_headers if header != "Employee ID"
    ]
    _write_employee_workbook(source_path, headers)

    with pytest.raises(IngestionError, match="missing required columns: Employee ID"):
        read_source(source_path, EMPLOYEES)


def test_rejects_duplicate_required_column(tmp_path: Path) -> None:
    source_path = tmp_path / EMPLOYEES.filename
    headers = [*EMPLOYEES.source_headers, "Employee ID"]
    _write_employee_workbook(source_path, headers)

    with pytest.raises(IngestionError, match="duplicate required columns: Employee ID"):
        read_source(source_path, EMPLOYEES)


def test_rejects_unexpected_column(tmp_path: Path) -> None:
    source_path = tmp_path / EMPLOYEES.filename
    headers = [*EMPLOYEES.source_headers, "Unexpected Field"]
    _write_employee_workbook(source_path, headers)

    with pytest.raises(IngestionError, match="unexpected columns: Unexpected Field"):
        read_source(source_path, EMPLOYEES)


def test_rejects_missing_expected_sheet(tmp_path: Path) -> None:
    source_path = tmp_path / EMPLOYEES.filename
    _write_employee_workbook(
        source_path,
        list(EMPLOYEES.source_headers),
        sheet_name="Wrong Sheet",
    )

    with pytest.raises(IngestionError, match="Expected sheet 'Employee Master Data'"):
        read_source(source_path, EMPLOYEES)


def test_ingestion_loads_raw_tables_and_is_idempotent(tmp_path: Path) -> None:
    database_path = tmp_path / "staffing.duckdb"
    data_dir = REPOSITORY_ROOT / "data"

    first_counts = ingest(data_dir, database_path)
    second_counts = ingest(data_dir, database_path)

    assert first_counts == {"hr_employees": 25, "project_assignments": 26}
    assert second_counts == first_counts

    with duckdb.connect(str(database_path), read_only=True) as connection:
        assert connection.execute(
            "select count(*) from raw.hr_employees"
        ).fetchone() == (25,)
        assert connection.execute(
            "select count(*) from raw.project_assignments"
        ).fetchone() == (26,)
        assert connection.execute(
            "select project_code, weekly_hours from raw.project_assignments "
            "where assignment_id = 'ASGN-0001'"
        ).fetchone() == ("PROJ-2024-001", "36")
        for table_name, row_count in first_counts.items():
            assert connection.execute(
                f"select _source_snapshot_date, count(*) from raw.{table_name} "
                "group by _source_snapshot_date"
            ).fetchall() == [(date(2026, 2, 2), row_count)]


def test_missing_source_file_has_clear_error(tmp_path: Path) -> None:
    missing_spec = replace(EMPLOYEES, filename="missing.xlsx")

    with pytest.raises(IngestionError, match="Expected source file does not exist"):
        read_source(tmp_path / missing_spec.filename, missing_spec)


def test_invalid_refresh_preserves_both_raw_snapshots(tmp_path: Path) -> None:
    employee_path = tmp_path / EMPLOYEES.filename
    assignment_path = tmp_path / ASSIGNMENTS.filename
    database_path = tmp_path / "staffing.duckdb"
    _write_employee_workbook(employee_path, list(EMPLOYEES.source_headers))

    assignments = Workbook()
    assignments.active.title = ASSIGNMENTS.sheet_name
    assignments.active["A2"] = "Report Date: 02/02/2026"
    assignments.active.append([])
    assignments.active.append(ASSIGNMENTS.source_headers)
    assignments.active.append(
        ["ASGN-1", "EMP-1", "PROJ-1", "Test project", "Lead", "2024-02-01", 20, "Y"]
    )
    assignments.save(assignment_path)

    ingest(tmp_path, database_path)
    with duckdb.connect(str(database_path), read_only=True) as connection:
        previous_employees = connection.execute(
            "select * from raw.hr_employees order by employee_id"
        ).fetchall()
        previous_assignments = connection.execute(
            "select * from raw.project_assignments order by assignment_id"
        ).fetchall()

    employees = load_workbook(employee_path)
    employees.active["B5"] = "Changed first name"
    employees.save(employee_path)
    employees.close()
    assert read_source(employee_path, EMPLOYEES)[0]["first_name"] == "Changed first name"

    assignments.active["A4"] = None
    assignments.save(assignment_path)
    assignments.close()

    with pytest.raises(IngestionError, match="missing required columns: Assignment ID"):
        ingest(tmp_path, database_path)

    with duckdb.connect(str(database_path), read_only=True) as connection:
        assert connection.execute(
            "select * from raw.hr_employees order by employee_id"
        ).fetchall() == previous_employees
        assert connection.execute(
            "select * from raw.project_assignments order by assignment_id"
        ).fetchall() == previous_assignments
